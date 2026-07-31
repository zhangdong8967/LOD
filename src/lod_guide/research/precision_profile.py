"""
EP17-style precision-profile / LoQ engine (G3).

Given multi-pool replicates (concentration ≈ pool mean), estimate:
  - per-pool mean, SD, CV%
  - CV(%) vs concentration curve (log-linear or piecewise linear on log10 mean)
  - LoQ at target CV% by interpolation (lowest concentration where fitted CV <= target)
  - recommended n for CV-of-mean at LoQ-adjacent pools

Works on Analyse-it troponin Example 3 (tier C) and any {pool_id -> values} map.
"""

from __future__ import annotations

import json
from dataclasses import asdict, dataclass
from pathlib import Path

import numpy as np
from scipy import stats

from lod_guide.data.download_real import REMOTE_DATASETS, download_dataset
from lod_guide.s5_sample_size import sample_size_cv_precision

OUT = Path(__file__).resolve().parents[3] / "data" / "research"


@dataclass
class PoolPrecision:
    pool_id: str
    n: int
    mean: float
    sd: float
    cv_pct: float
    stratum: str = ""


@dataclass
class PrecisionProfileResult:
    pools: list[PoolPrecision]
    target_cv_pct: float
    loq_estimate: float | None
    loq_method: str
    fit: dict
    n_design_at_loq_adjacent: dict
    evidence_tier: str
    notes: list[str]


def pool_stats(values: np.ndarray, pool_id: str, stratum: str = "") -> PoolPrecision:
    x = np.asarray(values, dtype=float)
    x = x[np.isfinite(x)]
    mu = float(np.mean(x))
    sd = float(np.std(x, ddof=1)) if x.size > 1 else float("nan")
    cv = 100.0 * sd / mu if mu != 0 and np.isfinite(sd) else float("nan")
    return PoolPrecision(pool_id, int(x.size), mu, sd, cv, stratum)


def fit_cv_logmean(pools: list[PoolPrecision]) -> dict:
    """
    Fit CV% = a + b * log10(mean) by OLS on pools with mean>0 and finite CV.
    Declining CV with concentration is the usual precision-profile shape.
    """
    means, cvs = [], []
    for p in pools:
        if p.mean > 0 and np.isfinite(p.cv_pct):
            means.append(p.mean)
            cvs.append(p.cv_pct)
    means = np.asarray(means, dtype=float)
    cvs = np.asarray(cvs, dtype=float)
    if means.size < 2:
        return {"ok": False, "reason": "need >=2 pools"}
    x = np.log10(means)
    slope, intercept, r, pval, se = stats.linregress(x, cvs)
    return {
        "ok": True,
        "model": "CV_pct = a + b*log10(mean)",
        "a": float(intercept),
        "b": float(slope),
        "r": float(r),
        "pvalue": float(pval),
        "stderr_slope": float(se),
    }


def loq_from_profile(
    pools: list[PoolPrecision],
    *,
    target_cv_pct: float = 20.0,
    fit: dict | None = None,
) -> tuple[float | None, str]:
    """
    Estimate LoQ as the lowest concentration where CV% reaches target.

    Prefer: interpolate between adjacent empirical pools (crossing target).
    Fallback: invert log-linear fit if slope < 0.
    """
    ordered = sorted([p for p in pools if p.mean > 0 and np.isfinite(p.cv_pct)], key=lambda p: p.mean)
    if not ordered:
        return None, "no_valid_pools"

    # empirical crossing: find consecutive pools where CV goes from >target to <=target
    for i in range(len(ordered) - 1):
        a, b = ordered[i], ordered[i + 1]
        if a.cv_pct >= target_cv_pct >= b.cv_pct and a.cv_pct != b.cv_pct:
            # linear interpolate in log-concentration space
            w = (a.cv_pct - target_cv_pct) / (a.cv_pct - b.cv_pct)
            log_c = (1 - w) * np.log10(a.mean) + w * np.log10(b.mean)
            return float(10**log_c), "empirical_log_interpolation"
        if b.cv_pct >= target_cv_pct >= a.cv_pct and a.cv_pct != b.cv_pct:
            w = (b.cv_pct - target_cv_pct) / (b.cv_pct - a.cv_pct)
            log_c = (1 - w) * np.log10(b.mean) + w * np.log10(a.mean)
            return float(10**log_c), "empirical_log_interpolation"

    # all below or all above
    if all(p.cv_pct <= target_cv_pct for p in ordered):
        return float(ordered[0].mean), "all_pools_meet_target_use_lowest_mean"
    if all(p.cv_pct > target_cv_pct for p in ordered):
        fit = fit or fit_cv_logmean(ordered)
        if fit.get("ok") and fit["b"] < 0:
            # target = a + b log10(c) => log10(c) = (target-a)/b
            log_c = (target_cv_pct - fit["a"]) / fit["b"]
            return float(10**log_c), "extrapolated_loglinear_fit"
        return None, "all_pools_above_target_no_reliable_extrapolation"

    fit = fit or fit_cv_logmean(ordered)
    if fit.get("ok") and abs(fit["b"]) > 1e-12:
        log_c = (target_cv_pct - fit["a"]) / fit["b"]
        return float(10**log_c), "loglinear_fit_inversion"
    return None, "unable_to_estimate"


def load_troponin_pools() -> tuple[list[PoolPrecision], str]:
    from collections import defaultdict
    from openpyxl import load_workbook

    meta = next(d for d in REMOTE_DATASETS if d.dataset_id == "analyseit_ep17_troponin_precision")
    path = download_dataset(meta, force=False)
    wb = load_workbook(path, data_only=True, read_only=True)
    pools: list[PoolPrecision] = []
    for sheet_name in wb.sheetnames:
        if not sheet_name.lower().startswith("reagent"):
            continue
        rows = list(wb[sheet_name].iter_rows(values_only=True))
        header_i = None
        for i, r in enumerate(rows[:10]):
            if r and r[0] and str(r[0]).strip().upper() == "ID":
                header_i = i
                break
        if header_i is None:
            continue
        bucket: dict[str, list[float]] = defaultdict(list)
        for r in rows[header_i + 1 :]:
            if not r or r[0] is None or r[1] is None:
                continue
            try:
                bucket[str(r[0]).strip()].append(float(r[1]))
            except (TypeError, ValueError):
                continue
        for pid, vals in sorted(bucket.items()):
            pools.append(pool_stats(np.array(vals), pid, stratum=sheet_name))
    wb.close()
    return pools, "C"


def evaluate_precision_profile(
    pools: list[PoolPrecision] | None = None,
    *,
    target_cv_pct: float = 20.0,
    evidence_tier: str = "C",
) -> PrecisionProfileResult:
    notes: list[str] = []
    if pools is None:
        pools, evidence_tier = load_troponin_pools()
        notes.append("Loaded Analyse-it EP17 Example 3 troponin pools (tier C synthetic demo).")

    fit = fit_cv_logmean(pools)
    loq, method = loq_from_profile(pools, target_cv_pct=target_cv_pct, fit=fit)

    n_design = {}
    if loq is not None:
        # nearest pool at or above LoQ
        candidates = [p for p in pools if p.mean >= loq and np.isfinite(p.cv_pct) and p.mean > 0]
        if not candidates:
            candidates = sorted(pools, key=lambda p: abs(p.mean - loq))[:1]
        adj = min(candidates, key=lambda p: p.mean) if candidates else None
        if adj is not None and adj.mean > 0:
            som = abs(adj.sd / adj.mean)
            n_design = {
                "adjacent_pool": adj.pool_id,
                "adjacent_mean": adj.mean,
                "sigma_over_mu": som,
                "n_for_CV_mean_20": sample_size_cv_precision(20.0, som) if som > 0 else None,
                "n_for_CV_mean_10": sample_size_cv_precision(10.0, som) if som > 0 else None,
                "n_for_target_CV_mean": sample_size_cv_precision(target_cv_pct, som) if som > 0 else None,
            }

    if evidence_tier == "C":
        notes.append("Tier C: do not interpret LoQ as clinical truth — method demonstration.")

    return PrecisionProfileResult(
        pools=pools,
        target_cv_pct=target_cv_pct,
        loq_estimate=loq,
        loq_method=method,
        fit=fit,
        n_design_at_loq_adjacent=n_design,
        evidence_tier=evidence_tier,
        notes=notes,
    )


def run_precision_profile_study(*, target_cv_pct: float = 20.0) -> dict:
    OUT.mkdir(parents=True, exist_ok=True)
    res20 = evaluate_precision_profile(target_cv_pct=20.0)
    res10 = evaluate_precision_profile(pools=res20.pools, target_cv_pct=10.0, evidence_tier=res20.evidence_tier)
    payload = {
        "study_id": "PRECISION-PROFILE-v1",
        "target_20": {
            "loq": res20.loq_estimate,
            "method": res20.loq_method,
            "fit": res20.fit,
            "n_design": res20.n_design_at_loq_adjacent,
        },
        "target_10": {
            "loq": res10.loq_estimate,
            "method": res10.loq_method,
            "n_design": res10.n_design_at_loq_adjacent,
        },
        "pools": [asdict(p) for p in res20.pools],
        "evidence_tier": res20.evidence_tier,
        "notes": res20.notes,
    }
    path = OUT / "precision_profile_loq.json"
    path.write_text(json.dumps(payload, indent=2), encoding="utf-8")
    payload["path"] = str(path)
    return payload


if __name__ == "__main__":
    print(json.dumps(run_precision_profile_study(), indent=2)[:2000])
