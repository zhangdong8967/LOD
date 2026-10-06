"""Parsers for multi-assay public LoD datasets (qPCR, EP17 Excel, IC, summaries)."""

from __future__ import annotations

import csv
from pathlib import Path


def load_zenodo_qpcr_cq(path: Path) -> dict:
    rows = []
    with path.open(newline="", encoding="utf-8", errors="replace") as f:
        reader = csv.DictReader(f)
        for r in reader:
            rows.append(r)

    def parse_cq(v: str) -> float | None:
        v = (v or "").strip()
        if v in {"", "N/A", "NA", "Undetermined", "undetermined"}:
            return None
        try:
            return float(v)
        except ValueError:
            return None

    def parse_sq(v: str) -> float | None:
        v = (v or "").strip()
        if v in {"", "Undetermined", "N/A", "NA"}:
            return None
        try:
            return float(v)
        except ValueError:
            return None

    blanks_cq: list[float] = []
    blank_strata: list[str] = []
    low_cq: list[float] = []
    hits_by_sq: dict[float, list[int]] = {}

    for r in rows:
        sample = (r.get("Sample") or "").strip()
        cq = parse_cq(r.get("Cq") or "")
        sq = parse_sq(r.get("SQ") or "")
        well = (r.get("Well") or "A01").strip()
        stratum = well[0] if well else "A"

        is_ntc = sample.upper().startswith("NTC") or (sq is None and "NTC" in sample.upper())
        if is_ntc or sample.upper() == "NTC":
            blanks_cq.append(45.0 if cq is None else cq)
            blank_strata.append(stratum)
            continue

        if sq is None:
            continue
        hits_by_sq.setdefault(sq, []).append(1 if cq is not None else 0)
        if cq is not None and sq > 0:
            low_cq.append(cq)

    positive_sqs = sorted(s for s in hits_by_sq if s > 0)
    lowest_sq = positive_sqs[0] if positive_sqs else None
    low_level_cq: list[float] = []
    if lowest_sq is not None:
        for r in rows:
            sq = parse_sq(r.get("SQ") or "")
            cq = parse_cq(r.get("Cq") or "")
            if sq == lowest_sq and cq is not None:
                low_level_cq.append(cq)

    hit_summary = {
        float(sq): {"n": len(v), "hits": int(sum(v)), "rate": float(sum(v) / len(v)) if v else 0.0}
        for sq, v in sorted(hits_by_sq.items())
    }
    return {
        "blank_values": blanks_cq,
        "blank_strata": blank_strata,
        "low_level_values": low_level_cq or low_cq[:50],
        "hit_summary": hit_summary,
        "n_rows": len(rows),
        "path": str(path),
        "assay_family": "qpcr",
    }


def load_ep17_estradiol_xlsx(path: Path) -> dict:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    reagents: dict = {}
    for sheet_name in wb.sheetnames:
        if not sheet_name.lower().startswith("reagent"):
            continue
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        header_idx = None
        header = []
        for i, r in enumerate(rows[:15]):
            if r is None:
                continue
            cells = [str(h).strip() if h is not None else "" for h in r]
            joined = " ".join(cells).lower()
            if "estradiol" in joined or ("level" in joined and "sample" in joined):
                header_idx = i
                header = cells
                break
        if header_idx is None:
            continue
        val_col = None
        level_col = None
        for i, h in enumerate(header):
            hl = h.lower()
            if "estradiol" in hl or "pg/ml" in hl or "(pg" in hl:
                val_col = i
            if hl.startswith("level"):
                level_col = i if level_col is None else level_col
        if val_col is None:
            val_col = len(header) - 1
        if level_col is None:
            level_col = 0

        blanks: list[float] = []
        lows: list[float] = []
        for r in rows[header_idx + 1 :]:
            if r is None or val_col is None or val_col >= len(r):
                continue
            level = str(r[level_col] if level_col < len(r) else "").strip().lower()
            try:
                val = float(r[val_col])
            except (TypeError, ValueError):
                continue
            if "blank" in level or level in {"0", "lob"}:
                blanks.append(val)
            elif "low" in level or "lod" in level:
                lows.append(val)
            elif level.startswith("b"):
                blanks.append(val)
            elif level.startswith("l"):
                lows.append(val)

        if blanks or lows:
            reagents[sheet_name] = {
                "blank_values": blanks,
                "low_level_values": lows,
                "n_blank": len(blanks),
                "n_low": len(lows),
            }
    wb.close()
    if not reagents:
        return {
            "assay_family": "immunoassay_estradiol",
            "path": str(path),
            "reagents": {},
            "blank_values": [],
            "low_level_values": [],
            "blank_strata": [],
            "n_rows": 0,
            "parser_note": "No blank/low rows parsed — inspect workbook layout",
        }
    # prefer sheet with most blanks
    primary_name = max(reagents, key=lambda k: reagents[k]["n_blank"])
    primary = reagents[primary_name]
    return {
        "assay_family": "immunoassay_estradiol",
        "path": str(path),
        "reagents": reagents,
        "primary_sheet": primary_name,
        "blank_values": primary["blank_values"],
        "low_level_values": primary["low_level_values"],
        "blank_strata": [0] * len(primary["blank_values"]),
        "n_rows": sum(r["n_blank"] + r["n_low"] for r in reagents.values()),
        "parser_note": "Blank vs Low classified by Level column text",
    }


def load_ep17_probit_xlsx(path: Path) -> dict:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    sheet = wb[wb.sheetnames[0]]
    rows = list(sheet.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return {"assay_family": "probit_hit", "path": str(path), "hit_summary": {}, "n_rows": 0}

    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    c_col = 0
    lot_cols: list[int] = []
    # Find true header row (may be after title lines)
    header_idx = 0
    for i, r in enumerate(rows[:15]):
        if r is None:
            continue
        cells = [str(h).strip() if h is not None else "" for h in r]
        joined = " ".join(cells).lower()
        if "cfl" in joined or ("lot" in joined and "c (" in joined):
            header_idx = i
            header = cells
            break
    for i, h in enumerate(header):
        hl = h.lower()
        if hl.startswith("c ") or "cfl" in hl or hl == "c":
            c_col = i
        if "lot" in hl:
            lot_cols.append(i)
    if not lot_cols:
        lot_cols = [i for i, h in enumerate(header) if "lot" in h.lower()]
    data_rows = rows[header_idx + 1 :]

    hits_by_c: dict[float, list[int]] = {}
    n_data = 0
    for r in data_rows:
        if r is None or c_col >= len(r) or r[c_col] is None:
            continue
        try:
            c = float(r[c_col])
        except (TypeError, ValueError):
            continue
        for lc in lot_cols:
            if lc >= len(r) or r[lc] is None:
                continue
            v = r[lc]
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                # sometimes R is replicate id; only treat 0/1 as detect if binary-like
                if float(v) in (0.0, 1.0):
                    hit = int(float(v))
                else:
                    continue
            else:
                s = str(v).strip().lower()
                if s in {"", "na", "n/a"}:
                    continue
                hit = 1 if s in {"1", "true", "detect", "detected", "pos", "positive", "yes"} else 0
            hits_by_c.setdefault(c, []).append(hit)
            n_data += 1

    hit_summary = {
        float(c): {"n": len(v), "hits": int(sum(v)), "rate": float(sum(v) / len(v)) if v else 0.0}
        for c, v in sorted(hits_by_c.items())
    }
    return {
        "assay_family": "probit_hit",
        "path": str(path),
        "hit_summary": hit_summary,
        "n_rows": n_data,
        "lot_columns": lot_cols,
    }


def load_ep17_troponin_xlsx(path: Path) -> dict:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    vals: list[float] = []
    pool_ids: list[str] = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        for r in rows[1:]:
            if not r:
                continue
            pool = str(r[0]) if r[0] is not None else ""
            for cell in r[1:]:
                try:
                    vals.append(float(cell))
                    pool_ids.append(pool)
                    break
                except (TypeError, ValueError):
                    continue
    wb.close()
    n_blankish = max(1, len(vals) // 9)  # lowest pool proxy
    return {
        "assay_family": "cardiac_precision_profile",
        "path": str(path),
        "blank_values": vals[:n_blankish],
        "low_level_values": vals,
        "blank_strata": pool_ids[:n_blankish],
        "n_rows": len(vals),
        "evidence_note": "Tier C synthetic EP17 precision-profile demo",
    }


def load_cation_ic_csv(path: Path) -> dict:
    return _load_ic_lod_csv(path, assay_family="analytical_chemistry_ic_cation")


def load_anion_ic_csv(path: Path) -> dict:
    return _load_ic_lod_csv(path, assay_family="analytical_chemistry_ic_anion")


def _load_ic_lod_csv(path: Path, *, assay_family: str) -> dict:
    text = path.read_text(encoding="utf-8", errors="replace")
    delim = ";" if text.count(";") >= text.count(",") else ","
    lines = [ln for ln in text.splitlines() if ln.strip()]
    reader = csv.DictReader(lines, delimiter=delim)
    rows = list(reader)
    numeric_cols: list[str] = []
    if rows:
        for k in rows[0].keys():
            if k is None:
                continue
            kl = k.lower()
            if "concentration" in kl or kl.endswith("_conc"):
                numeric_cols.append(k)
                continue
            for r in rows:
                v = r.get(k)
                if v is None or str(v).strip() == "":
                    continue
                try:
                    float(str(v).replace(",", "."))
                    numeric_cols.append(k)
                except ValueError:
                    pass
                break
    series = {c: [] for c in numeric_cols}
    for r in rows:
        for c in numeric_cols:
            try:
                series[c].append(float(str(r[c]).replace(",", ".")))
            except (TypeError, ValueError):
                pass
    preferred = [
        c
        for c in numeric_cols
        if any(x in c.lower() for x in ("magnesium", "sodium", "chloride", "fluoride", "nitrate", "sulfate"))
    ]
    primary = preferred[0] if preferred else (numeric_cols[0] if numeric_cols else None)
    values = series.get(primary, []) if primary else []
    return {
        "assay_family": assay_family,
        "path": str(path),
        "columns": numeric_cols,
        "series": series,
        "blank_values": values,
        "blank_strata": [0] * len(values),
        "low_level_values": values,
        "n_rows": len(rows),
        "evidence_note": "Small n IC LOD/LOQ replicates; diversity exemplar (tier B)",
    }


def load_carobene_lombardy_eqa_csv(path: Path) -> dict:
    """
    Carobene et al. multicenter clinical chemistry raw replicates (Mendeley 10.17632/r4szpw8sn5).
    Ten labs × (AST, ALT, ALP, LDH, CHE, CK, bilirubin, creatinine, calcium);
    M1–M3 control materials (6 reps) and C01–C15 serum pools (3 reps).
    Tier A for routine chemistry precision / multi-lab design-effect — not blank-LoB.
    """
    import re
    from collections import defaultdict

    text = path.read_text(encoding="latin-1")
    rows = list(csv.reader(text.splitlines(), delimiter=";"))
    if len(rows) < 3:
        return {"assay_family": "clinical_chemistry_multilab", "path": str(path), "n_rows": 0, "levels": []}

    lab_header = rows[0]
    analyte_header = rows[1]
    # lab blocks: 9 analytes each starting at col 2, 11, 20, ...
    analytes = ["AST", "ALT", "ALP", "LDH", "CHE", "CK", "Bili_Tot", "Creatinine", "Calcium"]
    lab_starts: list[tuple[str, int]] = []
    for i, cell in enumerate(lab_header):
        if cell and str(cell).strip().upper().startswith("LAB"):
            lab_starts.append((str(cell).strip(), i))
    if not lab_starts:
        # fallback fixed stride
        lab_starts = [(f"LAB_{k+1}", 2 + 9 * k) for k in range(10)]

    long_rows: list[dict] = []
    for r in rows[2:]:
        if not r or not r[0]:
            continue
        code = str(r[0]).strip()
        m = re.match(r"^(M\d+)\.(\d+)$", code)
        if m:
            material, rep = m.group(1), int(m.group(2))
            kind = "control_material"
        else:
            m2 = re.match(r"^(C\d+)\.(\d+)$", code)
            if m2:
                material, rep = m2.group(1), int(m2.group(2))
                kind = "serum_pool"
            else:
                material, rep, kind = code, 1, "other"
        for lab_name, start in lab_starts:
            for j, analyte in enumerate(analytes):
                idx = start + j
                if idx >= len(r):
                    continue
                raw = str(r[idx]).strip() if r[idx] is not None else ""
                if not raw:
                    continue
                try:
                    val = float(raw.replace(",", "."))
                except ValueError:
                    continue
                long_rows.append(
                    {
                        "material": material,
                        "replicate": rep,
                        "kind": kind,
                        "lab": lab_name,
                        "analyte": analyte,
                        "value": val,
                    }
                )

    # precision by material × lab × analyte
    groups: dict[tuple, list[float]] = defaultdict(list)
    for rec in long_rows:
        groups[(rec["material"], rec["lab"], rec["analyte"], rec["kind"])].append(rec["value"])

    levels = []
    for (material, lab, analyte, kind), vals in sorted(groups.items()):
        import numpy as np

        arr = np.asarray(vals, dtype=float)
        mu = float(arr.mean())
        sd = float(arr.std(ddof=1)) if arr.size > 1 else float("nan")
        cv = 100.0 * sd / mu if mu and np.isfinite(sd) else float("nan")
        levels.append(
            {
                "level": f"{material}|{lab}|{analyte}",
                "material": material,
                "lab": lab,
                "analyte": analyte,
                "kind": kind,
                "n": int(arr.size),
                "mean": mu,
                "sd": sd,
                "cv_pct": cv,
                "values": arr.tolist(),
            }
        )

    # representative continuous vector: creatinine across all labs/pools (common LoD-adjacent analyte)
    crea = [rec["value"] for rec in long_rows if rec["analyte"] == "Creatinine"]
    labs = sorted({r["lab"] for r in long_rows})
    return {
        "assay_family": "clinical_chemistry_multilab_eqa",
        "path": str(path),
        "n_labs": len(labs),
        "labs": labs,
        "analytes": analytes,
        "n_rows": len(long_rows),
        "levels": levels,
        "long_rows": long_rows,
        "blank_values": crea[: max(5, len(crea) // 10)] if crea else [],
        "low_level_values": crea,
        "blank_strata": [0] * (max(5, len(crea) // 10) if crea else 0),
        "evidence_note": (
            "Multicenter routine chemistry raw replicates (Carobene et al., Clin Chim Acta 2025; "
            "Mendeley 10.17632/r4szpw8sn5). Precision/multi-lab DE — not blank-LoB EP17 design."
        ),
    }


def load_cptac_msstats_lobd_csv(path: Path, *, focus_peptide: str = "LPPGLLANFTLLR") -> dict:
    """
    CPTAC Assay Portal spike-in calibration (JHU_DChan_HZhang_ZZhang) as shipped in
    Bioconductor MSstatsLOBD::raw_data (Galitzine et al., MCP 2018; Artistic-2.0).

    Preprocessing follows the MSstatsLOBD vignette: sum fragment peak areas per
    peptide × replicate × SampleGroup × concentration × file, then analyze blank
    (Concentration==0) light areas with EP17-style parametric/nonparametric LoB.

    Domain boundary: targeted proteomics blank–LoB case study — not a routine
    clinical chemistry analyzer panel.
    """
    from collections import defaultdict

    import numpy as np

    rows: list[dict] = []
    with path.open(newline="", encoding="utf-8", errors="replace") as f:
        reader = csv.DictReader(f)
        for r in reader:
            try:
                conc = float(r.get("Concentration") or "")
                light = float(r.get("light.Area") or "")
            except ValueError:
                continue
            heavy_raw = (r.get("heavy.Area") or "").strip()
            try:
                heavy = float(heavy_raw) if heavy_raw else float("nan")
            except ValueError:
                heavy = float("nan")
            rows.append(
                {
                    "peptide": (r.get("Peptide.Sequence") or "").strip(),
                    "concentration": conc,
                    "replicate": (r.get("Replicate") or "").strip(),
                    "sample_group": (r.get("SampleGroup") or "").strip(),
                    "file_name": (r.get("File.Name") or "").strip(),
                    "light": light,
                    "heavy": heavy,
                }
            )

    # Sum fragments → one intensity per peptide × file × conc × replicate × group
    sums: dict[tuple, list[float]] = defaultdict(lambda: [0.0, 0.0, 0])
    for rec in rows:
        key = (
            rec["peptide"],
            rec["replicate"],
            rec["sample_group"],
            rec["concentration"],
            rec["file_name"],
        )
        sums[key][0] += rec["light"]
        if np.isfinite(rec["heavy"]):
            sums[key][1] += float(rec["heavy"])
        sums[key][2] += 1

    agg: list[dict] = []
    for key, (a_light, a_heavy, n_frag) in sums.items():
        peptide, replicate, sample_group, concentration, file_name = key
        agg.append(
            {
                "peptide": peptide,
                "replicate": replicate,
                "sample_group": sample_group,
                "concentration": float(concentration),
                "file_name": file_name,
                "A_light": float(a_light),
                "A_heavy": float(a_heavy),
                "n_fragments": int(n_frag),
            }
        )

    peptides = sorted({a["peptide"] for a in agg if a["peptide"]})
    concentrations = sorted({a["concentration"] for a in agg})

    per_peptide: list[dict] = []
    for pep in peptides:
        blank = [a["A_light"] for a in agg if a["peptide"] == pep and a["concentration"] == 0.0]
        spiked = [a for a in agg if a["peptide"] == pep and a["concentration"] > 0.0]
        low_conc = min((a["concentration"] for a in spiked), default=None)
        low = [a["A_light"] for a in spiked if low_conc is not None and a["concentration"] == low_conc]
        b_arr = np.asarray(blank, dtype=float)
        l_arr = np.asarray(low, dtype=float)
        per_peptide.append(
            {
                "peptide": pep,
                "n_blank": int(b_arr.size),
                "n_low": int(l_arr.size),
                "blank_mean": float(b_arr.mean()) if b_arr.size else None,
                "blank_sd": float(b_arr.std(ddof=1)) if b_arr.size > 1 else None,
                "low_concentration": low_conc,
                "low_mean": float(l_arr.mean()) if l_arr.size else None,
                "blank_values": b_arr.tolist(),
                "low_level_values": l_arr.tolist(),
            }
        )

    focus = next((p for p in per_peptide if p["peptide"] == focus_peptide), None)
    if focus is None and per_peptide:
        focus = max(per_peptide, key=lambda p: p["n_blank"])
        focus_peptide = focus["peptide"]

    # Concentration grid (pooled mean of peptide light means) for precision-style summary
    levels = []
    for conc in concentrations:
        vals = [a["A_light"] for a in agg if a["concentration"] == conc]
        arr = np.asarray(vals, dtype=float)
        if arr.size == 0:
            continue
        mu = float(arr.mean())
        sd = float(arr.std(ddof=1)) if arr.size > 1 else float("nan")
        levels.append(
            {
                "level": f"conc={conc}",
                "concentration": conc,
                "n": int(arr.size),
                "mean": mu,
                "sd": sd,
                "cv_pct": (100.0 * sd / mu) if mu and np.isfinite(sd) else float("nan"),
                "values": arr.tolist(),
            }
        )

    blank_values = (focus or {}).get("blank_values") or []
    low_values = (focus or {}).get("low_level_values") or []
    return {
        "assay_family": "proteomics_cptac_msstats_lobd",
        "path": str(path),
        "n_rows": len(rows),
        "n_aggregated": len(agg),
        "n_peptides": len(peptides),
        "peptides": peptides,
        "concentrations": concentrations,
        "focus_peptide": focus_peptide,
        "aggregated": agg,
        "per_peptide": per_peptide,
        "levels": levels,
        "blank_values": blank_values,
        "low_level_values": low_values,
        "blank_strata": [0] * len(blank_values),
        "evidence_note": (
            "CPTAC / MSstatsLOBD calibration_raw (43 peptides; blank Concentration=0 + "
            "spike grid). Proteomics blank–LoB case study — not clinical chemistry analyzer LoB."
        ),
    }


def load_radioimmuno_lod_txt(path: Path) -> dict:
    lines = path.read_text(encoding="utf-8", errors="replace").strip().splitlines()
    records = []
    for line in lines:
        parts = [p.strip() for p in line.replace("\t", " ").split() if p.strip()]
        nums = []
        for p in parts:
            try:
                nums.append(float(p.replace(",", ".")))
            except ValueError:
                continue
        if nums:
            records.append({"raw": line, "nums": nums})
    return {
        "assay_family": "radioimmuno_summary",
        "path": str(path),
        "records": records,
        "n_rows": len(records),
        "evidence_note": "Summary blank mean/SD only — tier B evidence",
        "blank_values": [r["nums"][0] for r in records if r["nums"]],
    }


def load_plos_trec_lod_xlsx(path: Path) -> dict:
    """PLOS ONE S3 — TREC qPCR LoD-adjacent Cp replicates (Munich newborn screening)."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True)
    ws = wb["(3)"]
    rows = list(ws.iter_rows(values_only=True))[1:]
    low_cq: list[float] = []
    hits_by_label: dict[str, list[int]] = {}
    for label, cp, _conc in rows:
        if not label:
            continue
        key = str(label).strip()
        detected = isinstance(cp, (int, float))
        hits_by_label.setdefault(key, []).append(1 if detected else 0)
        if detected and "1:100" in key:
            low_cq.append(float(cp))
    hit_summary = {
        k: {"n": len(v), "hits": int(sum(v)), "rate": float(sum(v) / len(v)) if v else 0.0}
        for k, v in sorted(hits_by_label.items())
    }
    return {
        "assay_family": "hospital_qpcr_lod_replicates",
        "path": str(path),
        "hit_summary": hit_summary,
        "low_level_values": low_cq,
        "blank_values": low_cq[: max(5, len(low_cq) // 3)] if low_cq else [],
        "n_rows": len(rows),
        "evidence_note": "Hospital newborn-screening TREC LoD replicates (PLOS ONE 2024 S3; CC BY 4.0)",
    }


def parse_real_dataset(dataset_id: str, path: Path) -> dict:
    if dataset_id.startswith("zenodo_mus"):
        return load_zenodo_qpcr_cq(path)
    if dataset_id == "analyseit_ep17_estradiol":
        return load_ep17_estradiol_xlsx(path)
    if dataset_id == "analyseit_ep17_probit":
        return load_ep17_probit_xlsx(path)
    if dataset_id == "analyseit_ep17_troponin_precision":
        return load_ep17_troponin_xlsx(path)
    if dataset_id == "zenodo_cation_ic_lod":
        return load_cation_ic_csv(path)
    if dataset_id == "zenodo_anion_ic_lod":
        return load_anion_ic_csv(path)
    if dataset_id == "mendeley_carobene_lombardy_eqa":
        return load_carobene_lombardy_eqa_csv(path)
    if dataset_id == "bioc_msstats_lobd_cptac":
        return load_cptac_msstats_lobd_csv(path)
    if dataset_id == "zenodo_radioimmuno_lod_summary":
        return load_radioimmuno_lod_txt(path)
    if dataset_id == "zenodo_radioimmuno_precision":
        return load_radioimmuno_precision_txt(path)
    if dataset_id == "zenodo_radioimmuno_linearity":
        return load_radioimmuno_linearity_txt(path)
    if dataset_id == "zenodo_rmn_qualification":
        return load_rmn_qualification_xlsx(path)
    if dataset_id == "plos_bzdok2024_trec_lod":
        return load_plos_trec_lod_xlsx(path)
    raise KeyError(f"No parser for {dataset_id}")


def load_radioimmuno_precision_txt(path: Path) -> dict:
    """
    Zenodo 12664367 Precision.txt — per-sample replicate mean/SD/CV (Anal. Chem.).
    Not full raw DPM vectors, but sample-level precision suitable for CV profile.
    Tier A for precision-profile style analysis (published validation raw tables).
    """
    text = path.read_text(encoding="utf-8", errors="replace")
    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
    summary = {}
    data_start = None
    for i, ln in enumerate(lines):
        if ln.upper() == "DATA":
            data_start = i + 1
            break
        if "\t" in ln:
            k, v = ln.split("\t", 1)
            try:
                summary[k.strip()] = float(v)
            except ValueError:
                summary[k.strip()] = v.strip()
    rows = []
    if data_start is not None and data_start < len(lines):
        header = [h.strip() for h in lines[data_start].split("\t")]
        for ln in lines[data_start + 1 :]:
            parts = ln.split("\t")
            if len(parts) < 3:
                continue
            rec = dict(zip(header, parts))
            try:
                rows.append(
                    {
                        "sample_name": rec.get("sample_name", parts[0]),
                        "mean": float(rec.get("replicate_average_DPM", parts[1])),
                        "sd": float(rec.get("replicate_sd_DPM", parts[2])),
                        "cv_pct": float(rec.get("relative_sd_%", parts[3])) if len(parts) > 3 else float("nan"),
                    }
                )
            except ValueError:
                continue
    means = [r["mean"] for r in rows]
    return {
        "assay_family": "radioimmuno_precision_profile",
        "path": str(path),
        "summary": summary,
        "pool_rows": rows,
        "blank_values": means[: max(1, len(means) // 5)] if means else [],
        "low_level_values": means,
        "blank_strata": [0] * (max(1, len(means) // 5) if means else 0),
        "n_rows": len(rows),
        "evidence_note": "Published replicate-mean/SD table (Anal. Chem. / Zenodo 12664367)",
    }


def load_radioimmuno_linearity_txt(path: Path) -> dict:
    """Zenodo linearity file: concentration (ng) × DPM replicates."""
    text = path.read_text(encoding="utf-8", errors="replace")
    lines = [ln.rstrip("\r") for ln in text.splitlines()]
    data_i = None
    for i, ln in enumerate(lines):
        if ln.strip().upper() == "DATA":
            data_i = i + 1
            break
    rows = []
    if data_i is not None and data_i < len(lines):
        header = [h.strip() for h in lines[data_i].split("\t")]
        for ln in lines[data_i + 1 :]:
            if not ln.strip():
                continue
            parts = ln.split("\t")
            # sample_name, ng, DPM, ...
            if len(parts) < 3:
                continue
            try:
                ng = float(parts[1])
                dpm = float(parts[2])
            except ValueError:
                continue
            rows.append({"sample_name": parts[0].strip(), "ng": ng, "DPM": dpm})
    # group by concentration
    from collections import defaultdict

    by_c: dict[float, list[float]] = defaultdict(list)
    for r in rows:
        by_c[r["ng"]].append(r["DPM"])
    hit_or_prec = []
    for c, vals in sorted(by_c.items()):
        arr = vals
        mu = sum(arr) / len(arr)
        if len(arr) > 1:
            var = sum((x - mu) ** 2 for x in arr) / (len(arr) - 1)
            sd = var**0.5
            cv = 100 * sd / mu if mu else float("nan")
        else:
            sd, cv = float("nan"), float("nan")
        hit_or_prec.append({"concentration_ng": c, "n": len(arr), "mean_DPM": mu, "sd_DPM": sd, "cv_pct": cv, "values": arr})
    flat = [r["DPM"] for r in rows]
    return {
        "assay_family": "radioimmuno_linearity_continuous",
        "path": str(path),
        "rows": rows,
        "by_concentration": hit_or_prec,
        "blank_values": flat[: max(3, len(flat) // 4)] if flat else [],
        "low_level_values": flat,
        "blank_strata": [0] * (max(3, len(flat) // 4) if flat else 0),
        "n_rows": len(rows),
        "evidence_note": "Published linearity replicates ng×DPM (Zenodo 12664367)",
    }


def load_rmn_qualification_xlsx(path: Path) -> dict:
    """
    Zenodo 17108929 RMN assay qualification — replicate ID50 values by dilution.
    Continuous precision / CV across dilution levels (tier A, clinical immunoserology).
    """
    from openpyxl import load_workbook
    from collections import defaultdict

    wb = load_workbook(path, data_only=True, read_only=True)
    sheet_name = "Lin,CV,RE,Range" if "Lin,CV,RE,Range" in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return {"assay_family": "rmn_qualification", "path": str(path), "n_rows": 0, "levels": []}
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    # map columns
    idx = {h: i for i, h in enumerate(header)}
    r1 = idx.get("Replicate1_ID50")
    r2 = idx.get("Replicate2_ID50")
    dil = idx.get("log10_dilution")
    desc = idx.get("description")
    levels: dict[str, list[float]] = defaultdict(list)
    all_vals = []
    n = 0
    for r in rows[1:]:
        if r is None:
            continue
        try:
            v1 = float(r[r1]) if r1 is not None and r[r1] is not None else None
            v2 = float(r[r2]) if r2 is not None and r[r2] is not None else None
        except (TypeError, ValueError):
            continue
        key = None
        if dil is not None and r[dil] is not None:
            key = f"log10_dilution={r[dil]}"
        elif desc is not None:
            key = str(r[desc])
        else:
            key = "all"
        for v in (v1, v2):
            if v is not None:
                levels[key].append(v)
                all_vals.append(v)
                n += 1
    level_stats = []
    for k, vals in sorted(levels.items(), key=lambda kv: kv[0]):
        import numpy as np

        arr = np.asarray(vals, dtype=float)
        mu = float(arr.mean())
        sd = float(arr.std(ddof=1)) if arr.size > 1 else float("nan")
        cv = 100.0 * sd / mu if mu else float("nan")
        level_stats.append({"level": k, "n": int(arr.size), "mean": mu, "sd": sd, "cv_pct": cv, "values": arr.tolist()})
    return {
        "assay_family": "rmn_microneutralization_continuous",
        "path": str(path),
        "sheet": sheet_name,
        "levels": level_stats,
        "blank_values": all_vals[: max(5, len(all_vals) // 5)] if all_vals else [],
        "low_level_values": all_vals,
        "blank_strata": [0] * (max(5, len(all_vals) // 5) if all_vals else 0),
        "n_rows": n,
        "evidence_note": "Published RMN qualification replicates (Zenodo 17108929, CC-BY-4.0)",
    }
