"""
LoQ / precision main table (LoQ) for paper artifacts.

Provides:
  1) Analytic design rows: n for target CV% of the mean; n for CI half-width.
  2) Optional summary of Analyse-it EP17 troponin precision-profile pools (tier C).
"""

from __future__ import annotations

import csv
import json
from collections import defaultdict
from pathlib import Path

import numpy as np

from lod_guide.data.download_real import RAW_DIR, download_dataset, REMOTE_DATASETS
from lod_guide.s5_sample_size import sample_size_cv_precision, sample_size_mean_ci

OUT = Path(__file__).resolve().parents[3] / "paper_artifacts" / "tables"


def design_rows_cv_and_ci(
    *,
    sigma_over_mu_grid: tuple[float, ...] = (0.05, 0.10, 0.15, 0.20, 0.25, 0.30),
    target_cv_pct_grid: tuple[float, ...] = (20.0, 15.0, 10.0),
    half_width_frac_grid: tuple[float, ...] = (0.10, 0.05),
    alpha: float = 0.05,
) -> list[dict]:
    """
    Main LoQ design table.

    CV-of-mean rule: n >= (100 * (σ/μ) / target_CV%)^2
    CI half-width: for μ=1 scale, sd=σ/μ, half-width = frac * μ → n from z*sd/hw.
    """
    rows: list[dict] = []
    for som in sigma_over_mu_grid:
        for tcv in target_cv_pct_grid:
            n_cv = sample_size_cv_precision(tcv, som)
            rows.append(
                {
                    "block": "CV_of_mean",
                    "sigma_over_mu": som,
                    "target_CV_pct": tcv,
                    "CI_half_width_frac": "",
                    "n_required": n_cv,
                    "alpha": alpha,
                    "formula": "n=ceil((100*sigma_over_mu/target_CV_pct)^2)",
                    "note": "LoQ precision claim at a fixed low concentration",
                }
            )
        for frac in half_width_frac_grid:
            # normalize μ=1 so sd = sigma_over_mu
            n_ci = sample_size_mean_ci(sd=som, half_width=frac, alpha=alpha)
            rows.append(
                {
                    "block": "mean_CI_half_width",
                    "sigma_over_mu": som,
                    "target_CV_pct": "",
                    "CI_half_width_frac": frac,
                    "n_required": n_ci,
                    "alpha": alpha,
                    "formula": "n=ceil((z_{1-a/2}*sd/half_width)^2), mu=1 scale",
                    "note": "Alternative LoQ precision statement",
                }
            )
    return rows


def summarize_troponin_pools() -> list[dict]:
    """Per-pool mean/SD/CV% from EP17 Example 3 if file present."""
    meta = next(d for d in REMOTE_DATASETS if d.dataset_id == "analyseit_ep17_troponin_precision")
    try:
        path = download_dataset(meta, force=False)
    except Exception as exc:  # noqa: BLE001
        return [{"error": str(exc), "evidence_tier": "C", "note": "download failed"}]

    try:
        from openpyxl import load_workbook
    except ImportError:
        return [{"error": "openpyxl missing", "evidence_tier": "C"}]

    wb = load_workbook(path, data_only=True, read_only=True)
    rows_out: list[dict] = []
    for sheet_name in wb.sheetnames:
        if not sheet_name.lower().startswith("reagent"):
            continue
        ws = wb[sheet_name]
        data = list(ws.iter_rows(values_only=True))
        # find header
        header_i = None
        for i, r in enumerate(data[:10]):
            if r and r[0] and str(r[0]).strip().upper() == "ID":
                header_i = i
                break
        if header_i is None:
            continue
        pools: dict[str, list[float]] = defaultdict(list)
        for r in data[header_i + 1 :]:
            if not r or r[0] is None or r[1] is None:
                continue
            try:
                pools[str(r[0]).strip()].append(float(r[1]))
            except (TypeError, ValueError):
                continue
        for pool, vals in sorted(pools.items(), key=lambda kv: kv[0]):
            arr = np.asarray(vals, dtype=float)
            if arr.size == 0:
                continue
            mu = float(arr.mean())
            sd = float(arr.std(ddof=1)) if arr.size > 1 else float("nan")
            cv = 100.0 * sd / mu if mu != 0 and np.isfinite(sd) else float("nan")
            # n for 20% and 10% CV-of-mean using observed CV as sigma_over_mu proxy
            som = abs(sd / mu) if mu != 0 else float("nan")
            rows_out.append(
                {
                    "block": "troponin_EP17_pool",
                    "reagent_sheet": sheet_name,
                    "pool": pool,
                    "n_obs": int(arr.size),
                    "mean": mu,
                    "sd": sd,
                    "CV_pct": cv,
                    "n_for_CV_mean_20pct": (
                        int(sample_size_cv_precision(20.0, som)) if np.isfinite(som) and som > 0 else ""
                    ),
                    "n_for_CV_mean_10pct": (
                        int(sample_size_cv_precision(10.0, som)) if np.isfinite(som) and som > 0 else ""
                    ),
                    "evidence_tier": "C",
                    "note": "Synthetic EP17 demo matching target Mean/SD — method walk-through only",
                }
            )
    wb.close()
    return rows_out


def build_loq_precision_tables(*, out_dir: Path | None = None) -> dict[str, Path]:
    out_dir = out_dir or OUT
    out_dir.mkdir(parents=True, exist_ok=True)

    design = design_rows_cv_and_ci()
    pools = summarize_troponin_pools()

    main_path = out_dir / "Table02_loq_precision_main.csv"
    with main_path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(
            f,
            fieldnames=[
                "block",
                "sigma_over_mu",
                "target_CV_pct",
                "CI_half_width_frac",
                "n_required",
                "alpha",
                "formula",
                "note",
            ],
        )
        w.writeheader()
        w.writerows(design)

    pool_path = out_dir / "Table02b_troponin_precision_pools.csv"
    if pools and "error" not in pools[0]:
        keys = list(pools[0].keys())
        with pool_path.open("w", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, fieldnames=keys)
            w.writeheader()
            w.writerows(pools)
    else:
        pool_path.write_text(
            "note,detail\n"
            f"troponin_summary_unavailable,{json.dumps(pools)}\n",
            encoding="utf-8",
        )

    meta = {
        "table": "Table 2",
        "main_csv": str(main_path),
        "pool_csv": str(pool_path),
        "n_design_rows": len(design),
        "n_pool_rows": 0 if pools and "error" in pools[0] else len(pools),
        "interpretation": (
            "Use design rows for LoQ sample-size planning; troponin pools are tier-C "
            "illustrations of precision-profile structure, not clinical truth."
        ),
    }
    meta_path = out_dir / "Table02_loq_precision_meta.json"
    meta_path.write_text(json.dumps(meta, indent=2), encoding="utf-8")
    return {"main": main_path, "pools": pool_path, "meta": meta_path}
