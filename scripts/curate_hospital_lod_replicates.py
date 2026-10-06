#!/usr/bin/env python3
"""Export open-access hospital/clinical-lab LoD replicate panels to CSV."""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data" / "hospital_lod_replicates"
XLSX = ROOT / "data" / "raw" / "plos_bzdok2024_trec_lod_s3.xlsx"

PANELS: tuple[tuple[str, str, str, str, str, str], ...] = (
    (
        "plos_bzdok2024_trec_cxce_1to100",
        "TREC newborn screening qPCR",
        "TREC-PB CXCE 1:100 dilution (~1–3 copies/µl)",
        "TREC-PB CXCE 1:100",
        (
            "Limit of detection raw data for TREC compared for CXCE-buffer extractions "
            "(S3 Dataset sheet 3; per-replicate Cp)."
        ),
        "Dr von Hauner Children's Hospital / LMU Munich newborn screening laboratory",
    ),
    (
        "plos_bzdok2024_trec_magnapure_1to100",
        "TREC newborn screening qPCR",
        "TREC-PB MagnaPure 1:100 dilution (~1–3 copies/µl)",
        "TREC-PB MagnaPure 1:100",
        (
            "Limit of detection raw data for TREC compared for MagNA Pure 96 extractions "
            "(S3 Dataset sheet 3; per-replicate Cp)."
        ),
        "Dr von Hauner Children's Hospital / LMU Munich newborn screening laboratory",
    ),
    (
        "plos_bzdok2024_trec_cxce_1to1000",
        "TREC newborn screening qPCR",
        "TREC-PB CXCE 1:1000 dilution (sub-LoD panel)",
        "TREC-PB CXCE 1:1000",
        (
            "Sub-LoD dilution panel at 1:1000 with undetermined Cp recorded as non-detect "
            "(S3 Dataset sheet 3)."
        ),
        "Dr von Hauner Children's Hospital / LMU Munich newborn screening laboratory",
    ),
)

QUOTE_SOURCE = (
    "Bzdok J et al. PLOS ONE 2024; S3 Dataset (10.1371/journal.pone.0306329.s003); CC BY 4.0"
)
DOI = "10.1371/journal.pone.0306329"


def _detected(cp) -> int:
    if isinstance(cp, (int, float)):
        return 1
    return 0


def export_trec_panels() -> list[Path]:
    if not XLSX.exists():
        raise FileNotFoundError(f"Missing source xlsx: {XLSX}")
    OUT.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(XLSX, read_only=True)
    ws = wb["(3)"]
    rows = list(ws.iter_rows(values_only=True))[1:]
    written: list[Path] = []
    fields = [
        "source_id",
        "pathogen",
        "c_star",
        "replicate_id",
        "detected",
        "cq",
        "dilution_label",
        "protocol_quote",
        "quote_source",
        "citation_doi",
        "institution",
    ]
    for source_id, pathogen, c_star, dilution_prefix, quote, institution in PANELS:
        subset = [r for r in rows if r[0] and str(r[0]).strip() == dilution_prefix]
        path = OUT / f"{source_id}.csv"
        with path.open("w", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, fieldnames=fields)
            w.writeheader()
            for i, (label, cp, _conc) in enumerate(subset, start=1):
                w.writerow(
                    {
                        "source_id": source_id,
                        "pathogen": pathogen,
                        "c_star": c_star,
                        "replicate_id": i,
                        "detected": _detected(cp),
                        "cq": cp if isinstance(cp, (int, float)) else "",
                        "dilution_label": str(label).strip(),
                        "protocol_quote": quote,
                        "quote_source": QUOTE_SOURCE,
                        "citation_doi": DOI,
                        "institution": institution,
                    }
                )
        written.append(path)
    return written


if __name__ == "__main__":
    paths = export_trec_panels()
    for p in paths:
        print("wrote", p)
